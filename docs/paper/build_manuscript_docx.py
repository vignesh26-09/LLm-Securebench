"""Build an editable IEEE-style DOCX from the checked LaTeX manuscript.

Uses the repository's existing A4 conference artifact as the style/section
reference. It never executes model inference or changes research evidence.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
TEX = ROOT / "SecureLLMBench_IEEE_Manuscript.tex"
DOCX = ROOT / "SecureLLMBench_IEEE_Manuscript.docx"
TEMPLATE = ROOT.parents[1] / "artifacts" / "SecureLLMBench_IEEE_Conference_A4.docx"

REFERENCE_IDS = {
    "jbb": 1, "harmbench": 2, "agentdojo": 3,
    "injecagent": 4, "lostmiddle": 5, "strongreject": 6,
}
LABELS = {
    "tab:related": "Table I", "tab:checker": "Table II",
    "tab:config": "Table III", "tab:calibration": "Table IV",
    "tab:coverage": "Table V", "tab:risk": "Table VI",
    "tab:timing": "Table VII", "tab:comparison": "Table VIII",
    "fig:architecture": "Fig. 1", "fig:arms": "Fig. 2",
    "fig:workflow": "Fig. 3", "fig:risks": "Fig. 4",
    "fig:distance": "Fig. 5", "fig:models": "Fig. 6",
    "sec:existing-evidence": "Section IX-A", "sec:results": "Section IX",
}
EQUATIONS = {
    "eq:outcomes": "Y_final,j = 1{trial secret occurs in designated final response};  Y_earlier,j = max over eligible earlier responses of the same exact event;  Y_ever,j = max(Y_final,j, Y_earlier,j).",
    "eq:estimand": "Δ_m(d) = P(Y_final = 1 | A, d, m) − P(Y_final = 1 | T, d, m),   d ∈ {0, 2, 5}.",
    "eq:goal": "D_im(d) = mean(Y_im^A(d)) − mean(Y_im^T(d));   estimated Δ_m(d) = (1/G) Σ_i D_im(d),   G = 20 goals.",
    "eq:rc": "RC-AUC_raw = 1 − mean_t[ρ(x_t, b)] / δ_a;   RC_terminal = 1 − ρ(x_T, b) / δ_a.",
}


def clean(text: str) -> str:
    text = text.strip().replace("---", "—").replace("--", "–")
    # Keep inline mathematical operators legible in the editable Word version.
    # The TeX manuscript remains the authoritative equation source.
    for _ in range(4):
        updated = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
        if updated == text:
            break
        text = updated
    text = re.sub(r"\\sqrt\{([^{}]+)\}", r"√(\1)", text)
    for old, new in [(r"\cos", "cos"), (r"\min", "min"), (r"\max", "max"),
                     (r"\bar", "mean"), (r"\widehat", "estimated "),
                     (r"\mid", " | "), (r"\#", "#"), (r"\le", "≤"),
                     (r"\ge", "≥")]:
        text = text.replace(old, new)
    text = re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + ", ".join(str(REFERENCE_IDS[k]) for k in m.group(1).split(",")) + "]", text)
    text = re.sub(r"\\ref\{([^}]+)\}", lambda m: LABELS.get(m.group(1), m.group(1)), text)
    text = re.sub(r"\\label\{[^}]+\}", "", text)
    text = text.replace(r"\pending", "NOT YET MEASURED")
    text = text.replace(r"\insuff", "INSUFFICIENT VALIDATED EVIDENCE")
    text = text.replace(r"\texttt{/api/chat}", "/api/chat")
    text = re.sub(r"\\(?:textbf|textsc|texttt|emph|mathrm|operatorname|mathbf)\{([^{}]*)\}", r"\1", text)
    for old, new in [(r"\Delta", "Δ"), (r"\delta", "δ"), (r"\rho", "ρ"), (r"\kappa", "κ"),
                     (r"\ne", "≠"), (r"\in", "∈"), (r"\subseteq", "⊆"),
                     (r"\sum", "Σ"), (r"\prod", "∏"), (r"\times", "×"),
                     (r"\%", "%"), (r"\_", "_"), (r"\&", "&"),
                     (r"\!", ""), (r"\,", " "), (r"\quad", " "),
                     (r"\qquad", " "), (r"\rm", ""), (r"\text", "")]:
        text = text.replace(old, new)
    text = re.sub(r"\\url\{([^}]+)\}", r"\1", text)
    text = text.replace("$", "").replace("~", " ").replace("\\'", "'")
    text = text.replace(r"\`", "")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace("\\", "")
    text = text.replace("Table Table ", "Table ").replace("Sec. Section ", "Sec. ").replace("Fig. Fig. ", "Fig. ")
    return re.sub(r"\s+", " ", text).strip()


def set_columns(section, n: int) -> None:
    existing = section._sectPr.xpath("./w:cols")
    cols = existing[0] if existing else OxmlElement("w:cols")
    if not existing:
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "340")


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    current = tbl_pr.find(qn("w:tblBorders"))
    if current is not None:
        tbl_pr.remove(current)
    root = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        root.append(item)
    tbl_pr.append(root)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def setup(doc) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.59)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(2.54)
    set_columns(sec, 1)
    if "Title" not in doc.styles:
        doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    for style_name in ("Normal", "Body Text", "Title", "Heading 1", "Heading 2", "Abstract", "Keywords"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, bold, align in (("Heading 1", 10, True, WD_ALIGN_PARAGRAPH.CENTER),
                                    ("Heading 2", 9.5, True, WD_ALIGN_PARAGRAPH.LEFT)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = align


def caption(doc, text: str, *, figure: bool, number: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if figure else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = not figure
    romans = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII")
    run = p.add_run((f"Fig. {number}. " if figure else f"TABLE {romans[number - 1]}. ") + clean(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def add_figure(doc, block: str, number: int) -> None:
    match = re.search(r"\\caption\{(.*?)\}\s*\\label", block, re.S)
    title = match.group(1) if match else "Figure content pending."
    if number == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "Figure1_Backend_Architecture.png"), width=Inches(3.24))
    else:
        placeholder = re.search(r"\\placeholder\{([^{}]*)\}\{(.*?)\}\s*\\caption", block, re.S)
        text = clean((placeholder.group(1) + ". " + placeholder.group(2)) if placeholder else "Experimental result figure pending real observations.")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.right_indent = Cm(0.2)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(8)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F4F6")
        ppr.append(shd)
        borders_p = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "5")
            edge.set(qn("w:color"), "D9D9D9")
            borders_p.append(edge)
        ppr.append(borders_p)
    caption(doc, title, figure=True, number=number)


def add_table(doc, block: str, number: int, two_columns: bool) -> None:
    match = re.search(r"\\caption\{(.*?)\}\s*\\label", block, re.S)
    title = match.group(1) if match else "Table"
    wide = block.lstrip().startswith(r"\begin{table*}")
    caption(doc, title, figure=False, number=number)
    inside = re.search(r"\\begin\{tabular\}\{(?:[^{}]|\{[^{}]*\})*\}(.*?)\\end\{tabular\}", block, re.S)
    if not inside:
        return
    raw = re.sub(r"\\(?:toprule|midrule|bottomrule)", "", inside.group(1))
    raw_lines = [line.strip() for line in re.split(r"\\\\\s*", raw) if line.strip()]
    rows = [[clean(cell) for cell in line.split("&")] for line in raw_lines]
    ncol = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    width_cm = 17.75 if wide else 8.35
    widths = ([1.35] + [(width_cm - 1.35) / (ncol - 1)] * (ncol - 1)) if ncol > 1 else [width_cm]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Cm(width)
    for row_idx, row in enumerate(rows):
        for col_idx in range(ncol):
            cell = table.cell(row_idx, col_idx)
            cell.width = Cm(widths[col_idx])
            cell.text = row[col_idx] if col_idx < len(row) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                shade(cell, "E9EFF2")
            elif row_idx % 2 == 0:
                shade(cell, "F8FAFB")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = row_idx < len(rows) - 1
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(6.8 if wide else 7.0)
                    run.bold = row_idx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def blocks(source: str):
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line in (r"\maketitle",):
            i += 1
            continue
        match = re.match(r"\\(section|subsection)\{(.*?)\}(?:\\label\{[^}]*\})?$", line)
        if match:
            yield match.group(1), match.group(2)
            i += 1
            continue
        env = re.match(r"\\begin\{(figure|table\*?|equation|thebibliography)\}", line)
        if env:
            name = env.group(1)
            chunk = [line]
            i += 1
            while i < len(lines) and f"\\end{{{name}}}" not in chunk[-1]:
                chunk.append(lines[i])
                i += 1
            yield name, "\n".join(chunk)
            continue
        if line.startswith("\\"):
            i += 1
            continue
        chunk = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"\\(?:begin|section|subsection)", lines[i].lstrip()):
            chunk.append(lines[i].strip())
            i += 1
        yield "paragraph", " ".join(chunk)


def build() -> None:
    source = TEX.read_text(encoding="utf-8")
    title = re.search(r"\\title\{(.*?)\}\s*\\author", source, re.S).group(1)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S).group(1)
    keywords = re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", source, re.S).group(1)
    body = source.split(r"\end{IEEEkeywords}", 1)[1].split(r"\end{document}", 1)[0]
    doc = Document(TEMPLATE)
    setup(doc)
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(clean(title)).font.size = Pt(17)
    p = doc.add_paragraph("[AUTHOR NAME]  |  [DEPARTMENT]  |  [INSTITUTION]  |  [LOCATION]  |  [EMAIL]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    for run in p.runs:
        run.font.size = Pt(9)
    p = doc.add_paragraph()
    p.add_run("Abstract—").bold = True
    p.add_run(clean(abstract))
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph()
    p.add_run("Index Terms—").bold = True
    p.add_run(clean(keywords))
    p.paragraph_format.space_after = Pt(7)
    sec = doc.add_section(WD_SECTION_START.CONTINUOUS)
    set_columns(sec, 2)

    section_no = table_no = figure_no = 0
    for kind, value in blocks(body):
        if kind == "section":
            section_no += 1
            doc.add_paragraph(clean(value).upper(), style="Heading 1")
        elif kind == "subsection":
            doc.add_paragraph(clean(value), style="Heading 2")
        elif kind == "paragraph":
            p = doc.add_paragraph(clean(value))
            p.paragraph_format.widow_control = True
        elif kind == "equation":
            label = re.search(r"\\label\{([^}]+)\}", value)
            text = EQUATIONS.get(label.group(1), clean(value)) if label else clean(value)
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif kind.startswith("table"):
            table_no += 1
            if kind == "table*":
                set_columns(doc.add_section(WD_SECTION_START.CONTINUOUS), 1)
            add_table(doc, value, table_no, two_columns=True)
            if kind == "table*":
                set_columns(doc.add_section(WD_SECTION_START.CONTINUOUS), 2)
        elif kind == "figure":
            figure_no += 1
            add_figure(doc, value, figure_no)
        elif kind == "thebibliography":
            doc.add_paragraph("REFERENCES", style="Heading 1")
            for key, entry in re.findall(r"\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})", value, re.S):
                p = doc.add_paragraph(f"[{REFERENCE_IDS[key]}] {clean(entry)}")
                p.paragraph_format.left_indent = Cm(0.4)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.save(DOCX)
    print(f"saved {DOCX}; sections={section_no}, tables={table_no}, figures={figure_no}")


if __name__ == "__main__":
    build()
